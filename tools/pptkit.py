"""Shared pieces of the deck builders (tools/make_ppt.py, tools/make_ppt10.py): template helpers, colours, images and every measured number.

Keeps the template's logos, banner, cover photo, title style and Thank-you slide; replaces its instruction text with content,
adds a second EXTRA slide (prototype, rover). Images come from the Webots run (out/) and the team's rover render.
Render to check: tools/render_ppt.sh <pptx> [--pdf]
"""
import copy, glob, json, os, re
from lxml import etree
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
P = lambda *a: os.path.join(ROOT, *a)
IMG = dict(hall=P("out/run/cctv_004.00.jpg"), crowd=P("out/run/chase_012.00.jpg"), face=P("out/run/front_022.02.jpg"),
           dash=P("out/media/stills/dashboard_007.01.jpg"), cam=P("ppt/robot_cam_169.jpg"),
           rover=P("ppt/team_rover.jpg"), closeup=P("ppt/rover_sim_closeup.jpg"))
INK, TEAL, SAFFRON, MUTED = RGBColor(0x1B, 0x1F, 0x1E), RGBColor(0x0F, 0x5C, 0x52), RGBColor(0xB8, 0x55, 0x0D), RGBColor(0x55, 0x5E, 0x5B)
CARD, LINE, WHITE, MINT = RGBColor(0xF2, 0xF5, 0xF2), RGBColor(0xD3, 0xDB, 0xD6), RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0xDD, 0xEE, 0xE9)
FONT = "Montserrat"


def still(run, t):
    """The dashboard still of a run nearest to time t."""
    fs = glob.glob(P("out", "media", run, "stills", "dashboard_*.jpg"))
    return min(fs, key=lambda f: abs(float(re.search(r"_(\d+\.\d+)\.jpg", f).group(1)) - t))


# Every number on the results slides comes from these files, written by tools/bench.py, tools/run_webots.sh and tools/make_ab_media.py.
BENCH = json.load(open(P("out", "bench", "summary.json")))
AB = {k: json.load(open(P("out", f"ab_{k}", "summary.json"))) for k in ("baseline", "saathi", "vision")}
BB = json.load(open(P("out", "media", "ab", "blackbox.json")))
OB = {k: json.load(open(P("out", f"ob_{k}", "summary.json"))) for k in ("maponly", "vision")}
FIND, FALL = json.load(open(P("out", "find_dustbin", "summary.json"))), json.load(open(P("out", "fall_vision", "summary.json")))
FEAT = json.load(open(P("out", "media", "features", "facts.json")))
FALL_T = json.loads(open(P("out", "fall_vision", "fall_truth.jsonl")).readline())
VLM = json.load(open(P("out", "vlm", "answers.json")))


def _calib():
    """Drive test (tools/run_webots.sh calib --calib=1): 2 s forward, 2 s left, 2 s turning, all asked at 0.3 m/s or 0.6 rad/s."""
    import math
    m = json.load(open(P("out", "calib", "calib.json")))
    at = lambda s: min(m, key=lambda r: abs(r[0] - s))
    (t0, x0, y0, h0, _), (t1, x1, y1, _, _), (t2, x2, y2, _, _), (t3, _, _, h3, _) = at(0), at(2), at(4), at(6)
    fwd = ((x1 - x0) * math.cos(h0) + (y1 - y0) * math.sin(h0)) / (0.3 * (t1 - t0))
    side = (-(x2 - x1) * math.sin(h0) + (y2 - y1) * math.cos(h0)) / (0.3 * (t2 - t1))
    turn = (h3 - h0) / (0.6 * (t3 - t2))
    return max(abs(1 - fwd), abs(1 - side), abs(1 - turn))


CAL_ERR = _calib()
IMG.update(chart=P("out/bench/ab_chart.png"), traj=P("out/media/ab/trajectories.png"), side_base=P("out/media/ab/side_baseline_raw.jpg"),
           side_saathi=P("out/media/ab/side_saathi_raw.jpg"), ob_map=P("out/media/features/obstacle_ob_maponly.jpg"),
           ob_vis=P("out/media/features/obstacle_ob_vision.jpg"), found_box=P("out/media/features/found_box.jpg"),
           dustbin=P("out/media/features/dustbin.jpg"), arrive=P("out/media/features/arrive.jpg"), fall=P("out/media/features/fall.jpg"),
           stay=P("out/media/features/stay.jpg"), vision=still("ab_vision", 5.0), dash=still("ab_saathi", 6.5), blackbox=P("out/media/ab/blackbox.jpg"))
b0, s0 = BENCH["baseline"], BENCH["saathi"]
NEAR_DROP = round(100 * (1 - s0["near_misses_per_trip"] / b0["near_misses_per_trip"]))
H, BODY, SMALL, CAP = 18, 19, 16, 14  # header, body, card/bullet, caption sizes (pt) on the template's 20 x 11.25 in canvas


def text(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, align=PP_ALIGN.LEFT):
    """paras: dicts {runs: [(text, bold, color)], size, space_before, ls} or (text, size, bold, color) tuples."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap, tf.vertical_anchor = True, anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, p in enumerate(paras):
        if isinstance(p, tuple):
            p = dict(runs=[(p[0], p[2], p[3])], size=p[1])
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.line_spacing = p.get("ls", 1.12)
        if p.get("space_before"):
            para.space_before = Pt(p["space_before"])
        for t, bold, color in p["runs"]:
            r = para.add_run()
            r.text = t
            r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(p["size"]), bold, color
    return box


def para(s, size=BODY, color=INK, bold=False, space=4):
    return dict(runs=[(s, bold, color)], size=size, space_before=space)


def bullets(items, size=SMALL, color=INK, gap=5):
    return [dict(runs=[("•  ", True, TEAL), (s, False, color)], size=size, space_before=gap) for s in items]


def card(slide, x, y, w, h, fill=CARD, line=LINE, radius=0.06):
    s = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.adjustments[0] = radius
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb, s.line.width = line, Pt(1)
    s.shadow.inherit = False
    return s


def picture(slide, path, x, y, w, h, fit="cover", border=True):
    iw, ih = Image.open(path).size
    if fit == "cover":
        pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w), Inches(h))
        if iw / ih > w / h:
            pic.crop_left = pic.crop_right = (1 - (w / h) / (iw / ih)) / 2
        else:
            pic.crop_top = pic.crop_bottom = (1 - (iw / ih) / (w / h)) / 2
    else:
        s = min(w / iw, h / ih)
        pic = slide.shapes.add_picture(path, Inches(x + (w - iw * s) / 2), Inches(y + (h - ih * s) / 2), Inches(iw * s), Inches(ih * s))
    if border:
        pic.line.color.rgb, pic.line.width = LINE, Pt(0.75)
    return pic


def arrow(slide, x1, y1, x2, y2, color=MUTED, width=2.0):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb, c.line.width = color, Pt(width)
    e = etree.SubElement(c.line._get_or_add_ln(), qn("a:tailEnd"))
    e.set("type", "triangle"); e.set("w", "med"); e.set("len", "med")
    return c


def node(slide, x, y, w, h, title, body, fill=CARD, title_color=TEAL, line=LINE):
    card(slide, x, y, w, h, fill, line)
    text(slide, x + 0.18, y + 0.12, w - 0.36, h - 0.24, [(title, 16, True, title_color), para(body, 15)])


def remove(shape):
    shape._element.getparent().remove(shape._element)


def set_title(slide, title):
    """The template's title box keeps its look (Montserrat bold, black) and sits at one fixed place on every slide."""
    box = next(s for s in slide.shapes if s.has_text_frame and s.shape_type == 17)
    box.left, box.top, box.width, box.height = Inches(0.6), Inches(1.55), Inches(18.8), Inches(0.62)
    tf = box.text_frame
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    p = tf.paragraphs[0]
    for r in list(p.runs)[1:]:
        r._r.getparent().remove(r._r)
    r = p.runs[0]
    r.text = title
    r.font.name, r.font.size, r.font.bold, r.font.color.rgb = FONT, Pt(28), True, INK


def clear_body(slide):
    for s in list(slide.shapes):
        if s.shape_type == 6:  # the template's instruction groups
            remove(s)


def duplicate(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for s in list(new.shapes):
        remove(s)
    for s in src.shapes:
        el = copy.deepcopy(s._element)
        for blip in el.iter(qn("a:blip")):
            part = src.part.related_part(blip.get(qn("r:embed")))
            blip.set(qn("r:embed"), new.part.relate_to(part, RT.IMAGE))
        new.shapes._spTree.insert_element_before(el, "p:extLst")
    return new


def footer(slide, s):
    text(slide, 0.6, 10.62, 18.8, 0.4, [(s, 12, False, MUTED)])


